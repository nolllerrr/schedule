"""Bounded, value-only XLSX reading without a workbook-sized object graph."""

from __future__ import annotations

import posixpath
import struct
from dataclasses import dataclass, fields
from pathlib import Path
from typing import BinaryIO
from xml.etree.ElementTree import Element, fromstring, iterparse
from zipfile import ZipFile

from openpyxl.cell.text import Text
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format, is_timedelta_format
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.utils.datetime import (
    CALENDAR_MAC_1904,
    CALENDAR_WINDOWS_1900,
    from_excel,
    from_ISO8601,
)


class WorkbookLimitError(ValueError):
    """An XLSX is invalid or exceeds the supported resource budget."""


@dataclass(frozen=True, slots=True)
class WorkbookLimits:
    max_archive_bytes: int = 20 * 1024 * 1024
    max_members: int = 256
    max_uncompressed_bytes: int = 128 * 1024 * 1024
    max_part_bytes: int = 32 * 1024 * 1024
    max_metadata_bytes: int = 1024 * 1024
    max_sheets: int = 32
    max_rows: int = 4096
    max_columns: int = 512
    max_cells: int = 2_000_000
    max_value_cells: int = 100_000
    max_shared_strings: int = 100_000
    max_text_chars: int = 2_000_000
    max_styles: int = 200_000
    max_merges: int = 20_000

    def __post_init__(self):
        if any(
            not isinstance(getattr(self, field.name), int)
            or getattr(self, field.name) < 1
            for field in fields(self)
        ):
            raise ValueError("XLSX resource limits must be positive integers")


@dataclass(frozen=True, slots=True)
class ValueCell:
    value: object
    row: int
    column: int

    @property
    def coordinate(self) -> str:
        return f"{get_column_letter(self.column)}{self.row}"


@dataclass(slots=True)
class ValueWorksheet:
    title: str
    values: dict[tuple[int, int], object]
    max_row: int = 1
    max_column: int = 1

    def cell(self, row: int, column: int) -> ValueCell:
        # Missing cells must never become retained entries.
        return ValueCell(self.values.get((row, column)), row, column)

    def __getitem__(self, coordinate: str) -> ValueCell:
        return self.cell(*coordinate_to_tuple(coordinate))


class _XMLInput:
    def __init__(self, stream: BinaryIO, limit: int):
        self.stream, self.limit, self.read_bytes = stream, limit, 0
        self.tail = b""

    def read(self, size: int = -1) -> bytes:
        chunk = self.stream.read(min(size if size >= 0 else 65536, 65536))
        self.read_bytes += len(chunk)
        if self.read_bytes > self.limit:
            raise WorkbookLimitError("XLSX XML exceeds its size limit")
        scan = (self.tail + chunk).replace(b"\x00", b"")
        if b"<!DOCTYPE" in scan or b"<!ENTITY" in scan:
            raise WorkbookLimitError("XLSX XML declarations are not supported")
        self.tail = (self.tail + chunk)[-64:]
        return chunk


def _name(element: Element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def _records(stream: BinaryIO, limit: int, tags: set[str], *, row_starts: bool = False):
    """Discard completed XML nodes, retaining at most one requested record."""
    stack: list[Element] = []
    record_depth: int | None = None
    record_nodes = 0
    for event, node in iterparse(_XMLInput(stream, limit), events=("start", "end")):
        name = _name(node)
        if event == "start":
            stack.append(node)
            if len(stack) > 64:
                raise WorkbookLimitError("XLSX XML is nested too deeply")
            if row_starts and name == "row":
                yield "row_start", node
            if record_depth is None and name in tags:
                record_depth = len(stack)
                record_nodes = 0
            if record_depth is not None:
                record_nodes += 1
                if record_nodes > 4096:
                    raise WorkbookLimitError(
                        "An XLSX value contains too many XML nodes"
                    )
            continue
        if record_depth == len(stack):
            yield name, node
            record_depth = None
        if record_depth is None:
            if len(stack) > 1:
                stack[-2].remove(node)
            node.clear()
        stack.pop()


class WorkbookValues:
    def __init__(self, path: Path, limits: WorkbookLimits):
        self.path, self.limits = path, limits
        self.archive: ZipFile | None = None
        self.shared_strings: list[str] = []
        self.date_styles: set[int] = set()
        self.timedelta_styles: set[int] = set()
        self.epoch = CALENDAR_WINDOWS_1900
        self.sheets: list[tuple[str, str]] = []
        self.cells_read = self.text_chars = 0

    def __enter__(self):
        try:
            # Check the central-directory budget before ZipFile allocates a
            # ZipInfo object for every member (including in a malformed file).
            with self.path.open("rb") as source:
                source.seek(0, 2)
                size = source.tell()
                if size > self.limits.max_archive_bytes:
                    raise WorkbookLimitError("Compressed XLSX exceeds its size limit")
                source.seek(max(0, size - 65557))
                tail = source.read(65557)
            offset = tail.rfind(b"PK\x05\x06")
            if offset < 0 or len(tail) - offset < 22:
                raise WorkbookLimitError("XLSX ZIP directory is invalid")
            directory = struct.unpack_from("<4s4H2LH", tail, offset)
            if (
                directory[4] > self.limits.max_members
                or directory[5] > self.limits.max_metadata_bytes
            ):
                raise WorkbookLimitError("XLSX ZIP directory exceeds its size limit")
            self.archive = ZipFile(self.path)
            entries = self.archive.infolist()
            if len(entries) > self.limits.max_members or len(
                {p.filename for p in entries}
            ) != len(entries):
                raise WorkbookLimitError(
                    "XLSX contains too many or duplicate ZIP members"
                )
            if sum(p.file_size for p in entries) > self.limits.max_uncompressed_bytes:
                raise WorkbookLimitError("Uncompressed XLSX exceeds its size limit")
            if any(p.file_size > self.limits.max_part_bytes for p in entries):
                raise WorkbookLimitError("An XLSX part exceeds its size limit")
            workbook = self._metadata("xl/workbook.xml")
            relationships = {
                r.get("Id"): (r.get("Target", ""), r.get("Type", ""))
                for r in self._metadata("xl/_rels/workbook.xml.rels")
                if r.get("TargetMode") != "External"
            }
            for node in workbook.iter():
                if _name(node) == "workbookPr" and node.get(
                    "date1904", "0"
                ).lower() in {"1", "true"}:
                    self.epoch = CALENDAR_MAC_1904
                if _name(node) != "sheet":
                    continue
                relationship = next(
                    (v for k, v in node.attrib.items() if k.rsplit("}", 1)[-1] == "id"),
                    None,
                )
                target, relation_type = relationships.get(relationship, ("", ""))
                if relation_type.endswith("/chartsheet"):
                    continue
                part = posixpath.normpath(
                    target.lstrip("/")
                    if target.startswith("/")
                    else posixpath.join("xl", target)
                )
                if not part.startswith("xl/") or part not in self.archive.namelist():
                    raise WorkbookLimitError("XLSX worksheet relationship is invalid")
                self.sheets.append((node.get("name", ""), part))
            if len(self.sheets) > self.limits.max_sheets:
                raise WorkbookLimitError("XLSX contains too many sheets")
            self._read_strings()
            self._read_styles()
            return self
        except BaseException:
            self.__exit__(None, None, None)
            raise

    def __exit__(self, *exc):
        if self.archive is not None:
            self.archive.close()
        self.shared_strings.clear()
        self.date_styles.clear()
        self.timedelta_styles.clear()

    def _metadata(self, part: str) -> Element:
        assert self.archive is not None
        if self.archive.getinfo(part).file_size > self.limits.max_metadata_bytes:
            raise WorkbookLimitError("XLSX metadata exceeds its size limit")
        with self.archive.open(part) as source:
            reader = _XMLInput(source, self.limits.max_metadata_bytes)
            chunks = []
            while chunk := reader.read(65536):
                chunks.append(chunk)
            return fromstring(b"".join(chunks))

    def _text(self, text: str) -> str:
        self.text_chars += len(text)
        if len(text) > 32767 or self.text_chars > self.limits.max_text_chars:
            raise WorkbookLimitError("XLSX text exceeds its size limit")
        return text

    def _read_strings(self):
        assert self.archive is not None
        if "xl/sharedStrings.xml" not in self.archive.namelist():
            return
        with self.archive.open("xl/sharedStrings.xml") as source:
            for _, node in _records(source, self.limits.max_part_bytes, {"si"}):
                if len(self.shared_strings) >= self.limits.max_shared_strings:
                    raise WorkbookLimitError("XLSX contains too many shared strings")
                self.shared_strings.append(
                    self._text(Text.from_tree(node).content.replace("x005F_", ""))
                )

    def _read_styles(self):
        assert self.archive is not None
        if "xl/styles.xml" not in self.archive.namelist():
            return
        custom: dict[int, str] = {}
        stack: list[str] = []
        elements: list[Element] = []
        style_id = 0
        with self.archive.open("xl/styles.xml") as source:
            for event, node in iterparse(
                _XMLInput(source, self.limits.max_part_bytes), events=("start", "end")
            ):
                name = _name(node)
                if event == "start":
                    stack.append(name)
                    elements.append(node)
                    if len(stack) > 64:
                        raise WorkbookLimitError("XLSX XML is nested too deeply")
                    continue
                if name == "numFmt":
                    if len(custom) >= self.limits.max_styles:
                        raise WorkbookLimitError(
                            "XLSX contains too many number formats"
                        )
                    custom[int(node.get("numFmtId"))] = self._text(
                        node.get("formatCode", "")
                    )
                if name == "xf" and len(stack) > 1 and stack[-2] == "cellXfs":
                    if style_id >= self.limits.max_styles:
                        raise WorkbookLimitError("XLSX contains too many styles")
                    fmt_id = int(node.get("numFmtId", "0"))
                    fmt = custom.get(fmt_id, BUILTIN_FORMATS.get(fmt_id, ""))
                    if is_date_format(fmt):
                        self.date_styles.add(style_id)
                    if is_timedelta_format(fmt):
                        self.timedelta_styles.add(style_id)
                    style_id += 1
                if len(elements) > 1:
                    elements[-2].remove(node)
                node.clear()
                stack.pop()
                elements.pop()

    def _bounds(self, row: int, column: int):
        if (
            not 1 <= row <= self.limits.max_rows
            or not 1 <= column <= self.limits.max_columns
        ):
            raise WorkbookLimitError(
                "XLSX worksheet dimensions exceed the supported layout"
            )

    def read_sheet(self, title: str, part: str) -> ValueWorksheet:
        assert self.archive is not None
        sheet = ValueWorksheet(title, {})
        merges = []
        row = column = merge_area = 0
        with self.archive.open(part) as source:
            for kind, node in _records(
                source,
                self.limits.max_part_bytes,
                {"c", "dimension", "mergeCell"},
                row_starts=True,
            ):
                if kind == "row_start":
                    row, column = int(node.get("r", row + 1)), 0
                    self._bounds(row, 1)
                    sheet.max_row = max(sheet.max_row, row)
                elif kind == "dimension":
                    bounds = range_boundaries(node.get("ref", "A1"))
                    self._bounds(bounds[3], bounds[2])
                elif kind == "mergeCell":
                    bounds = range_boundaries(node.get("ref", ""))
                    self._bounds(bounds[1], bounds[0])
                    self._bounds(bounds[3], bounds[2])
                    merge_area += (bounds[2] - bounds[0] + 1) * (
                        bounds[3] - bounds[1] + 1
                    )
                    if (
                        len(merges) >= self.limits.max_merges
                        or merge_area > self.limits.max_cells
                    ):
                        raise WorkbookLimitError(
                            "XLSX merged ranges exceed their limit"
                        )
                    merges.append(bounds)
                    sheet.max_row = max(sheet.max_row, bounds[3])
                    sheet.max_column = max(sheet.max_column, bounds[2])
                else:
                    coordinate = node.get("r")
                    cell_row, column = (
                        coordinate_to_tuple(coordinate)
                        if coordinate
                        else (row, column + 1)
                    )
                    self._bounds(cell_row, column)
                    self.cells_read += 1
                    if self.cells_read > self.limits.max_cells:
                        raise WorkbookLimitError("XLSX contains too many cells")
                    sheet.max_row = max(sheet.max_row, cell_row)
                    sheet.max_column = max(sheet.max_column, column)
                    value = self._value(node)
                    if value is not None:
                        sheet.values[cell_row, column] = value
                        if len(sheet.values) > self.limits.max_value_cells:
                            raise WorkbookLimitError(
                                "XLSX contains too many value cells"
                            )
        # Excel/openpyxl discard values hidden inside a merged region.
        for first_col, first_row, last_col, last_row in merges:
            for r in range(first_row, last_row + 1):
                for c in range(first_col, last_col + 1):
                    if (r, c) != (first_row, first_col):
                        sheet.values.pop((r, c), None)
        return sheet

    def _value(self, node: Element):
        kind = node.get("t", "n")
        child = next(
            (c for c in node if _name(c) == ("is" if kind == "inlineStr" else "v")),
            None,
        )
        if child is None:
            return None  # An uncached formula has no value in data_only mode.
        if kind == "inlineStr":
            return self._text(Text.from_tree(child).content)
        value = child.text or None
        if value is None:
            return None
        if kind == "s":
            index = int(value)
            if not 0 <= index < len(self.shared_strings):
                raise WorkbookLimitError("XLSX shared-string reference is invalid")
            return self.shared_strings[index]
        if kind == "b":
            return bool(int(value))
        if kind == "d":
            return from_ISO8601(value)
        if kind == "n":
            number = float(value) if any(c in value for c in ".eE") else int(value)
            style = int(node.get("s", "0"))
            if style in self.date_styles:
                try:
                    return from_excel(
                        number, self.epoch, timedelta=style in self.timedelta_styles
                    )
                except (ValueError, OverflowError):
                    return "#VALUE!"
            return number
        return self._text(value)
