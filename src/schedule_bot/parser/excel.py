from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from schedule_bot.domain import Lesson, ParsedSchedule


class ScheduleParseError(ValueError):
    """Raised when a workbook does not match the supported schedule layout."""


RUSSIAN_MONTHS = {
    "января": 1,
    "февраля": 2,
    "марта": 3,
    "апреля": 4,
    "мая": 5,
    "июня": 6,
    "июля": 7,
    "августа": 8,
    "сентября": 9,
    "октября": 10,
    "ноября": 11,
    "декабря": 12,
}

WEEKDAYS = {
    0: "понедельник",
    1: "вторник",
    2: "среда",
    3: "четверг",
    4: "пятница",
    5: "суббота",
    6: "воскресенье",
}

WEEKDAY_ALIASES = {
    "понедельник": "понедельник",
    "вторник": "вторник",
    "вторинк": "вторник",
    "среда": "среда",
    "четверг": "четверг",
    "пятница": "пятница",
    "суббота": "суббота",
    "воскресенье": "воскресенье",
}

GROUP_PATTERN = re.compile(r"^[А-ЯЁA-Z]{1,4}\s*-\s*\d{1,3}$", re.IGNORECASE)
LESSON_PATTERN = re.compile(r"^(?P<number>\d+)\s*пара$", re.IGNORECASE)
CLASS_HOUR_PATTERN = re.compile(r"^кл\.?\s*час$", re.IGNORECASE)
TIME_PATTERN = re.compile(
    r"(?P<start>\d{1,2}[.:]\d{2})\s*[-–—]\s*(?P<end>\d{1,2}[.:]\d{2})"
)


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def normalize_group(value: object) -> str:
    text = normalize_text(value).upper()
    return re.sub(r"\s*-\s*", "-", text)


def parse_filename_date_range(filename: str) -> tuple[date, date]:
    name = Path(filename).stem
    full = re.search(
        r"(?P<sd>\d{1,2})\.(?P<sm>\d{1,2})\.(?P<sy>\d{2,4})"
        r"\s*-\s*"
        r"(?P<ed>\d{1,2})\.(?P<em>\d{1,2})\.(?P<ey>\d{2,4})",
        name,
    )
    if full:
        sy = _expand_year(int(full.group("sy")))
        ey = _expand_year(int(full.group("ey")))
        return (
            date(sy, int(full.group("sm")), int(full.group("sd"))),
            date(ey, int(full.group("em")), int(full.group("ed"))),
        )

    compact = re.search(
        r"(?P<sd>\d{1,2})\.\s*-\s*(?P<ed>\d{1,2})\."
        r"(?P<month>\d{1,2})\.(?P<year>\d{2,4})",
        name,
    )
    if compact:
        year = _expand_year(int(compact.group("year")))
        month = int(compact.group("month"))
        return (
            date(year, month, int(compact.group("sd"))),
            date(year, month, int(compact.group("ed"))),
        )
    raise ScheduleParseError(
        f"Cannot determine the schedule date range from filename: {filename}"
    )


def parse_sheet_date(title: str) -> date | None:
    normalized = normalize_text(title).lower()
    match = re.search(
        r"(?P<day>\d{1,2})\s+(?P<month>[а-яё]+)\s+(?P<year>\d{4})",
        normalized,
    )
    if match and match.group("month") in RUSSIAN_MONTHS:
        return date(
            int(match.group("year")),
            RUSSIAN_MONTHS[match.group("month")],
            int(match.group("day")),
        )
    numeric = re.search(
        r"(?P<day>\d{1,2})\.(?P<month>\d{1,2})\.(?P<year>\d{2,4})",
        normalized,
    )
    if numeric:
        return date(
            _expand_year(int(numeric.group("year"))),
            int(numeric.group("month")),
            int(numeric.group("day")),
        )
    return None


def _expand_year(value: int) -> int:
    return value + 2000 if value < 100 else value


class ExcelScheduleParser:
    def parse(
        self, path: str | Path, *, source_filename: str | None = None
    ) -> ParsedSchedule:
        source_path = Path(path)
        filename = source_filename or source_path.name
        start_date, end_date = parse_filename_date_range(filename)
        workbook = load_workbook(source_path, data_only=True, read_only=False)
        lessons: list[Lesson] = []
        parsed_dates: list[date] = []
        warnings: list[str] = []

        try:
            for worksheet in workbook.worksheets:
                sheet_date = parse_sheet_date(worksheet.title)
                if sheet_date is None:
                    warnings.append(
                        f"Лист {worksheet.title!r} пропущен: дата не распознана"
                    )
                    continue
                if not start_date <= sheet_date <= end_date:
                    continue

                group_columns = self._find_group_columns(worksheet)
                if not group_columns:
                    warnings.append(
                        f"Лист {worksheet.title!r} пропущен: группы не найдены"
                    )
                    continue

                day_row = self._find_matching_day_row(
                    worksheet, sheet_date, group_columns
                )
                if day_row is None:
                    warnings.append(
                        f"Лист {worksheet.title!r} пропущен: блок дня не найден"
                    )
                    continue

                parsed_dates.append(sheet_date)
                lessons.extend(
                    self._parse_day(
                        worksheet,
                        sheet_date=sheet_date,
                        day_row=day_row,
                        group_columns=group_columns,
                    )
                )
        finally:
            workbook.close()

        if not parsed_dates:
            raise ScheduleParseError("No schedule sheets were parsed")
        if not lessons:
            raise ScheduleParseError("The workbook contains no lessons")

        lessons.sort(
            key=lambda item: (item.lesson_date, item.group_name, item.position)
        )
        return ParsedSchedule(
            source_filename=filename,
            start_date=start_date,
            end_date=end_date,
            dates=tuple(sorted(set(parsed_dates))),
            lessons=tuple(lessons),
            warnings=tuple(warnings),
        )

    def _find_group_columns(self, worksheet: Worksheet) -> list[tuple[int, str]]:
        best: list[tuple[int, str]] = []
        for row in range(1, min(worksheet.max_row, 20) + 1):
            candidates: list[tuple[int, str]] = []
            for column in range(3, worksheet.max_column + 1):
                value = normalize_group(worksheet.cell(row, column).value)
                if value and GROUP_PATTERN.fullmatch(value):
                    candidates.append((column, value))
            if len(candidates) > len(best):
                best = candidates
        return best

    def _find_matching_day_row(
        self,
        worksheet: Worksheet,
        sheet_date: date,
        group_columns: list[tuple[int, str]],
    ) -> int | None:
        expected = WEEKDAYS[sheet_date.weekday()]
        candidates: list[int] = []
        for row in range(1, worksheet.max_row + 1):
            label = normalize_text(worksheet.cell(row, 1).value).lower()
            if WEEKDAY_ALIASES.get(label) == expected:
                candidates.append(row)
        if not candidates:
            return None
        return max(
            candidates,
            key=lambda row: self._day_content_score(
                worksheet, row=row, group_columns=group_columns
            ),
        )

    def _day_content_score(
        self,
        worksheet: Worksheet,
        *,
        row: int,
        group_columns: list[tuple[int, str]],
    ) -> int:
        score = 0
        for lesson_row in self._lesson_rows_after(worksheet, row):
            for start_column, _ in group_columns:
                if any(
                    normalize_text(worksheet.cell(r, c).value)
                    for r in range(lesson_row, lesson_row + 4)
                    for c in range(start_column, start_column + 4)
                ):
                    score += 1
        return score

    def _lesson_rows_after(
        self, worksheet: Worksheet, day_row: int
    ) -> Iterable[int]:
        stop_row = min(worksheet.max_row + 1, day_row + 33)
        for row in range(day_row + 1, stop_row):
            label = normalize_text(worksheet.cell(row, 1).value).lower()
            if label in WEEKDAY_ALIASES:
                stop_row = row
                break
        for row in range(day_row, stop_row):
            label = normalize_text(worksheet.cell(row, 2).value)
            if LESSON_PATTERN.fullmatch(label) or CLASS_HOUR_PATTERN.fullmatch(label):
                yield row

    def _parse_day(
        self,
        worksheet: Worksheet,
        *,
        sheet_date: date,
        day_row: int,
        group_columns: list[tuple[int, str]],
    ) -> list[Lesson]:
        lessons: list[Lesson] = []
        for position, lesson_row in enumerate(
            self._lesson_rows_after(worksheet, day_row), start=1
        ):
            lesson_label = normalize_text(worksheet.cell(lesson_row, 2).value)
            number_match = LESSON_PATTERN.fullmatch(lesson_label)
            lesson_number = int(number_match.group("number")) if number_match else None
            start_time, end_time = self._parse_time(
                worksheet.cell(lesson_row + 1, 2).value
            )

            for start_column, group_name in group_columns:
                raw_data = self._raw_block(
                    worksheet,
                    row=lesson_row,
                    column=start_column,
                )
                subject_parts = self._values(
                    worksheet,
                    rows=range(lesson_row, lesson_row + 2),
                    columns=range(start_column, start_column + 3),
                )
                teacher_parts = self._values(
                    worksheet,
                    rows=range(lesson_row + 2, lesson_row + 4),
                    columns=range(start_column, start_column + 3),
                )
                room_parts = self._values(
                    worksheet,
                    rows=range(lesson_row, lesson_row + 2),
                    columns=(start_column + 3,),
                )
                if not subject_parts:
                    continue

                subject = " / ".join(subject_parts)
                teacher = " / ".join(teacher_parts) or None
                room = self._format_room(room_parts)
                subgroup = self._detect_subgroup(subject)
                lessons.append(
                    Lesson(
                        lesson_date=sheet_date,
                        group_name=group_name,
                        position=position,
                        lesson_label=lesson_label,
                        lesson_number=lesson_number,
                        start_time=start_time,
                        end_time=end_time,
                        subject=subject,
                        teacher=teacher,
                        room=room,
                        subgroup=subgroup,
                        raw_data=raw_data,
                    )
                )
        return lessons

    @staticmethod
    def _parse_time(value: object) -> tuple[str | None, str | None]:
        match = TIME_PATTERN.search(normalize_text(value))
        if not match:
            return None, None
        return (
            ExcelScheduleParser._normalize_clock(match.group("start")),
            ExcelScheduleParser._normalize_clock(match.group("end")),
        )

    @staticmethod
    def _normalize_clock(value: str) -> str:
        hour, minute = re.split(r"[.:]", value)
        return f"{int(hour):02d}:{int(minute):02d}"

    @staticmethod
    def _values(
        worksheet: Worksheet,
        *,
        rows: Iterable[int],
        columns: Iterable[int],
    ) -> list[str]:
        result: list[str] = []
        for row in rows:
            for column in columns:
                value = normalize_text(worksheet.cell(row, column).value)
                if value and value not in result:
                    result.append(value)
        return result

    @staticmethod
    def _raw_block(
        worksheet: Worksheet, *, row: int, column: int
    ) -> dict[str, str]:
        result: dict[str, str] = {}
        for current_row in range(row, row + 4):
            for current_column in range(column, column + 4):
                cell = worksheet.cell(current_row, current_column)
                value = normalize_text(cell.value)
                if value:
                    result[cell.coordinate] = value
        return result

    @staticmethod
    def _format_room(parts: list[str]) -> str | None:
        if not parts:
            return None
        compact = "".join(part.replace(" ", "").lower() for part in parts)
        if compact in {"сп/з", "сп/зал", "спортзал"}:
            return "спортзал"
        if len(parts) == 2 and parts[0].endswith("/"):
            return parts[0] + parts[1]
        return " / ".join(parts)

    @staticmethod
    def _detect_subgroup(subject: str) -> str | None:
        found = sorted(
            set(
                re.findall(
                    r"\b([12])\s*п\s*/?\s*гр\.?",
                    subject,
                    flags=re.IGNORECASE,
                )
            )
        )
        return ",".join(found) or None
