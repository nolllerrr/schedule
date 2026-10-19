from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from schedule_bot.change_policy import changes_for_profile, should_notify
from schedule_bot.domain import Lesson, LessonSegment, ParsedSchedule, ScheduleChange
from schedule_bot.import_validation import validate_import
from schedule_bot.parser import ExcelScheduleParser
from schedule_bot.parser.workbook import WorkbookLimits, WorkbookValues
from schedule_bot.presentation import format_changes, format_schedule


DAY = date(2026, 10, 12)
PARTS = (
    LessonSegment("Математика 1 п/гр.", "Преподаватель А", "101", "1"),
    LessonSegment("Физика 2 п/гр.", "Преподаватель Б", "202", "2"),
)


def slot(parts=PARTS):
    return Lesson(
        DAY, "ИС-22", 1, "1 пара", 1, "09:00", "10:30", "Математика",
        teacher="Преподаватель А", room="101", raw_data={"C6": "диагностика"},
    ).with_segments(parts)


def workbook(tmp_path: Path, *, layout="vertical", shared=False, missing=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "12 октября 2026"
    for coordinate, value in {
        "C4": "ИС-22", "A6": "Понедельник", "B6": "1 пара", "B7": "09.00-10.30",
    }.items():
        sheet[coordinate] = value
    if layout == "vertical":
        sheet["C6"], sheet["C7"] = PARTS[0].subject, PARTS[1].subject
        sheet.merge_cells("C6:E6")
        sheet.merge_cells("C7:E7")
        sheet["C8"] = PARTS[0].teacher
        if shared:
            sheet.merge_cells("C8:E9")
        elif not missing:
            sheet["C9"] = PARTS[1].teacher
        sheet["F6"] = PARTS[0].room
        if shared:
            sheet.merge_cells("F6:F7")
        elif not missing:
            sheet["F7"] = PARTS[1].room
    elif layout == "horizontal":
        sheet["C6"], sheet["E6"] = PARTS[0].subject, PARTS[1].subject
        sheet.merge_cells("C6:D7")
        sheet.merge_cells("E6:E7")
        sheet["C8"] = PARTS[0].teacher
        if shared:
            sheet.merge_cells("C8:E9")
        else:
            sheet.merge_cells("C8:D9")
            if not missing:
                sheet["E8"] = PARTS[1].teacher
                sheet.merge_cells("E8:E9")
        sheet["F6"] = PARTS[0].room
        sheet.merge_cells("F6:F7")
    path = tmp_path / "12.10.2026.xlsx"
    return book, sheet, path


def parse_book(book, path):
    book.save(path)
    book.close()
    return ExcelScheduleParser().parse(path)


def test_legacy_serialization_and_single_explicit_segment_are_compatible():
    legacy = replace(slot(), segments=(), subject="Math", teacher="A", room="101", subgroup=None)
    explicit = legacy.with_segments((LessonSegment("Math", "A", "101"),))
    assert "segments" not in legacy.to_dict()
    assert explicit.to_dict()["segments"] == [
        {"subject": "Math", "teacher": "A", "room": "101", "subgroup": None}
    ]
    assert explicit.comparable_data() == legacy.comparable_data()
    assert explicit.raw_data == legacy.raw_data


def test_vertical_segments_keep_subject_teacher_room_mapping(tmp_path):
    book, _, path = workbook(tmp_path)
    parsed = parse_book(book, path)
    assert not parsed.validation_errors
    assert len(parsed.lessons) == 1
    assert parsed.lessons[0].segments == PARTS


@pytest.mark.parametrize("layout", ["vertical", "horizontal"])
def test_merged_common_values_are_shared_without_losing_subgroups(tmp_path, layout):
    book, _, path = workbook(tmp_path, layout=layout, shared=True)
    parsed = parse_book(book, path)
    assert not parsed.validation_errors
    parts = parsed.lessons[0].segments
    assert tuple(part.subgroup for part in parts) == ("1", "2")
    assert tuple(part.teacher for part in parts) == (PARTS[0].teacher,) * 2
    assert tuple(part.room for part in parts) == ("101", "101")


def test_horizontal_segments_follow_column_geometry(tmp_path):
    book, _, path = workbook(tmp_path, layout="horizontal")
    parsed = parse_book(book, path)
    assert not parsed.validation_errors
    assert parsed.lessons[0].segments == (PARTS[0], replace(PARTS[1], room="101"))


def test_missing_values_are_not_copied_from_other_vertical_subgroup(tmp_path):
    book, _, path = workbook(tmp_path, missing=True)
    parsed = parse_book(book, path)
    assert not parsed.validation_errors
    assert parsed.lessons[0].segments == (
        PARTS[0], replace(PARTS[1], teacher=None, room=None),
    )


def test_reader_preserves_merged_geometry_with_existing_budget(tmp_path):
    book, _, path = workbook(tmp_path)
    book.save(path)
    book.close()
    with WorkbookValues(path, WorkbookLimits()) as values:
        title, part = values.sheets[0]
        sheet = values.read_sheet(title, part)
        assert (3, 6, 5, 6) in sheet.merged_ranges
        assert sheet.cell_bounds(6, 4) == (3, 6, 5, 6)
        assert sheet.cell_bounds(8, 3) == (3, 8, 3, 8)


@pytest.mark.parametrize("case", ["diagonal", "one_cell", "horizontal_rooms", "extra_teacher"])
def test_ambiguous_parallel_geometry_becomes_structural_error(tmp_path, case):
    book, sheet, path = workbook(
        tmp_path, layout="horizontal" if case == "horizontal_rooms" else "vertical",
    )
    if case == "diagonal":
        sheet.unmerge_cells("C7:E7")
        sheet["C7"] = None
        sheet["E7"] = PARTS[1].subject
    elif case == "one_cell":
        sheet["C6"] = "Math 1 п/гр. / Physics 2 п/гр."
        sheet["C7"] = None
    elif case == "horizontal_rooms":
        sheet.unmerge_cells("F6:F7")
        sheet["F7"] = "202"
    else:
        sheet["E8"] = "Другой преподаватель"
    parsed = parse_book(book, path)
    assert parsed.validation_errors
    errors, _ = validate_import(parsed, [], {}, set())
    assert errors


def test_validation_rejects_empty_or_duplicate_segment_assignments():
    malformed = slot((PARTS[0], replace(PARTS[1], subject="", subgroup="1")))
    parsed = ParsedSchedule("12.10.2026.xlsx", DAY, DAY, (DAY,), (malformed,))
    errors, _ = validate_import(parsed, [], {}, set())
    assert any("отсутствует предмет" in error for error in errors)
    assert any("Повторяются подгруппы" in error for error in errors)


def test_teacher_projection_preserves_mapping_and_slot_count():
    projected = slot().for_teacher("Преподаватель А")
    assert projected is not None
    assert projected.segments == (PARTS[0],)
    assert projected.subject == PARTS[0].subject
    assert projected.room == "101"
    assert projected.raw_data == slot().raw_data
    assert slot().for_teacher("Нет такого") is None
    student = format_schedule([slot()], "ИС-22", DAY, "student")
    assert "<b>1 пара</b>" in student
    assert "Подгруппа 1: <b>Математика 1 п/гр.</b> • Преподаватель А • ауд. 101" in student
    assert "Подгруппа 2: <b>Физика 2 п/гр.</b> • Преподаватель Б • ауд. 202" in student
    teacher = format_schedule([projected], "Преподаватель А", DAY, "teacher")
    assert "Физика" not in teacher
    assert "202" not in teacher


def test_teacher_change_notifies_union_of_old_and_new_teachers_only():
    before = slot()
    after = slot((replace(PARTS[0], teacher="Преподаватель В"), PARTS[1]))
    changes = (ScheduleChange("modified", before, after),)
    assert should_notify(changes[0])
    old = changes_for_profile(changes, role="teacher", target="Преподаватель А")
    new = changes_for_profile(changes, role="teacher", target="Преподаватель В")
    unchanged = changes_for_profile(changes, role="teacher", target="Преподаватель Б")
    assert len(old) == len(new) == 1
    assert old[0].kind == "removed"
    assert new[0].kind == "added"
    assert not unchanged
    assert "Физика" not in format_changes(old)
    assert "Подгруппа 1 — Преподаватель" in format_changes(changes)


@pytest.mark.parametrize("kind", ["modified", "moved"])
def test_teacher_notification_keeps_real_slot_changes(kind):
    before = slot()
    after = replace(before, lesson_label="2 пара", lesson_number=2, position=2, start_time="11:00")
    changes = (ScheduleChange(kind, before, after),)
    selected = changes_for_profile(changes, role="teacher", target="Преподаватель Б")
    assert len(selected) == 1
    text = format_changes(selected)
    assert "1 пара → 2 пара" in text
    assert "09:00–10:30 → 11:00–10:30" in text
    if kind == "moved":
        assert "Перенесено" in text


def test_segment_cosmetic_subject_clarification_does_not_notify():
    before = slot((replace(PARTS[0], subject="МДК.01.01. / Разработка"),))
    after = before.with_segments((replace(before.segments[0], subject="МДК.01.01. КП / Разработка"),))
    assert not should_notify(ScheduleChange("modified", before, after))


def test_reordering_segment_cells_preserves_semantics():
    before = slot()
    after = replace(before, segments=tuple(reversed(before.segments)))
    assert before.comparable_data() == after.comparable_data()
    assert not should_notify(ScheduleChange("modified", before, after))


def test_segment_html_values_are_escaped():
    lesson = slot((replace(PARTS[0], subject="<Math>", teacher="A&B"), PARTS[1]))
    text = format_schedule([lesson], "ИС-22", DAY, "student")
    assert "&lt;Math&gt;" in text and "A&amp;B" in text
    assert "<Math>" not in text
