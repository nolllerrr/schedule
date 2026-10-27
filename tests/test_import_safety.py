from __future__ import annotations

import asyncio
import json
import os
import sqlite3
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from dataclasses import replace
from datetime import date, datetime, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest

pytestmark = pytest.mark.usefixtures("text_schedule_media")
from openpyxl import load_workbook

from schedule_bot.cli import _prepare_repository
from schedule_bot.domain import Lesson, ParsedSchedule
from schedule_bot.import_validation import ImportQuarantinedError
from schedule_bot.maintenance import backup_before_migration
from schedule_bot.notifier import ScheduleNotifier
from schedule_bot.parser import ExcelScheduleParser
from schedule_bot.repository import ScheduleRepository

DAY = date(2026, 9, 22)


def lesson(position: int = 1, *, subject: str = "Математика") -> Lesson:
    return Lesson(DAY, "ИС-22", position, f"{position} пара", position,
                  "09:00", "10:30", subject, "Иванов И.И.", "44")


def schedule(lessons: tuple[Lesson, ...]) -> ParsedSchedule:
    return ParsedSchedule("22.09.2026.xlsx", DAY, DAY, (DAY,), lessons,
                          date_groups={DAY: ("ИС-22",)})


def current(repository: ScheduleRepository) -> list[Lesson]:
    return repository.lessons_for(role="student", target="ИС-22", lesson_date=DAY)


@pytest.mark.parametrize("mutation", ["header", "blank_last_header", "slot", "day", "subject"])
def test_partial_workbook_does_not_replace_current_schedule(
    tmp_path: Path, schedule_workbook: Path, mutation: str,
) -> None:
    repository = ScheduleRepository(tmp_path / "schedule.db")
    parser = ExcelScheduleParser()
    original = parser.parse(schedule_workbook)
    repository.import_schedule(original, sha256="original")
    repository.save_profile(123, role="student", target="ИС-22")
    before = current(repository)
    workbook = load_workbook(schedule_workbook)
    sheet = workbook["22 сентября 2026"]
    if mutation == "header":
        sheet["K4"] = "НЕИЗВЕСТНАЯ ГРУППА"
    elif mutation == "blank_last_header":
        sheet["O4"] = None
    elif mutation == "slot":
        sheet["B106"] = "1 урок"
    elif mutation == "day":
        sheet["A106"] = "неизвестный день"
    elif mutation == "subject":
        sheet["K106"] = None
    workbook.save(schedule_workbook)
    workbook.close()
    partial = parser.parse(schedule_workbook)
    assert partial.validation_errors
    with pytest.raises(ImportQuarantinedError):
        repository.import_schedule(partial, sha256="partial")
    assert current(repository) == before
    assert repository.notification_outbox_counts() == {}
    with closing(sqlite3.connect(repository.database_path)) as connection, connection:
        assert connection.execute("SELECT COUNT(*) FROM imports").fetchone()[0] == 1
        row = connection.execute("SELECT reasons_json, snapshot_json FROM quarantined_imports").fetchone()
        assert json.loads(row[0])
        assert json.loads(row[1])["validation_errors"]
    # Retrying the same rejected SHA must not mark it as an imported file.
    with pytest.raises(ImportQuarantinedError):
        repository.import_schedule(partial, sha256="partial")
    with closing(sqlite3.connect(repository.database_path)) as connection, connection:
        assert connection.execute("SELECT COUNT(*) FROM quarantined_imports").fetchone()[0] == 1


def test_missing_group_preserves_even_previously_empty_group_coverage(tmp_path: Path) -> None:
    repository = ScheduleRepository(tmp_path / "schedule.db")
    original = replace(schedule((lesson(),)), date_groups={DAY: ("ИС-22", "ПД-12")})
    repository.import_schedule(original, sha256="original")
    with pytest.raises(ImportQuarantinedError, match="ПД-12"):
        repository.import_schedule(schedule((lesson(),)), sha256="missing-empty-group")
    assert current(repository) == [lesson()]


def test_missing_published_date_quarantines_entire_update(tmp_path: Path) -> None:
    repository = ScheduleRepository(tmp_path / "schedule.db")
    tomorrow = date(2026, 9, 23)
    original = replace(schedule((lesson(), replace(lesson(), lesson_date=tomorrow))),
                       end_date=tomorrow, dates=(DAY, tomorrow),
                       date_groups={DAY: ("ИС-22",), tomorrow: ("ИС-22",)})
    repository.import_schedule(original, sha256="original")
    partial = replace(original, dates=(DAY,), lessons=(lesson(subject="Физика"),),
                      date_groups={DAY: ("ИС-22",)})
    with pytest.raises(ImportQuarantinedError, match="даты"):
        repository.import_schedule(partial, sha256="partial")
    assert current(repository)[0].subject == "Математика"


def test_large_drop_needs_review_but_regular_cancellation_is_allowed(tmp_path: Path) -> None:
    repository = ScheduleRepository(tmp_path / "schedule.db")
    original = schedule(tuple(lesson(position) for position in range(1, 5)))
    repository.import_schedule(original, sha256="original")
    regular = replace(original, lessons=original.lessons[:3])
    result = repository.import_schedule(regular, sha256="one-cancelled")
    assert [change.kind for change in result.changes] == ["removed"]
    # Restore the full day so that 4 -> 1 triggers the anomaly threshold.
    repository.import_schedule(original, sha256="restored")
    drastic = replace(original, lessons=original.lessons[:1])
    with pytest.raises(ImportQuarantinedError, match="резко"):
        repository.import_schedule(drastic, sha256="drastic")
    assert len(current(repository)) == 4
    result = repository.import_schedule(drastic, sha256="drastic",
                                        approval_reason="Проверена отмена трёх пар в исходном файле")
    assert not result.skipped
    assert len(current(repository)) == 1
    with closing(sqlite3.connect(repository.database_path)) as connection, connection:
        assert connection.execute("SELECT status FROM quarantined_imports").fetchone()[0] == "approved"
        assert connection.execute("SELECT approval_reason FROM imports WHERE sha256='drastic'").fetchone()[0]


def test_structural_errors_cannot_be_approved(tmp_path: Path) -> None:
    repository = ScheduleRepository(tmp_path / "schedule.db")
    broken = replace(schedule((lesson(),)), validation_errors=("Не распознан заголовок",))
    with pytest.raises(ImportQuarantinedError):
        repository.import_schedule(broken, sha256="broken", approval_reason="Проверено")
    assert current(repository) == []


@pytest.mark.parametrize("broken", [
    replace(schedule((lesson(),)), lessons=(lesson(), lesson())),
    replace(schedule((lesson(),)), lessons=(replace(lesson(), lesson_date=date(2026, 9, 23)),)),
    replace(schedule((lesson(),)), confirmed_empty_dates=(DAY,)),
])
def test_invalid_normalized_data_never_activates(tmp_path: Path, broken: ParsedSchedule) -> None:
    repository = ScheduleRepository(tmp_path / "schedule.db")
    with pytest.raises(ImportQuarantinedError):
        repository.import_schedule(broken, sha256="broken")
    assert current(repository) == []


def test_explicit_empty_day_can_replace_schedule(tmp_path: Path, schedule_workbook: Path) -> None:
    repository = ScheduleRepository(tmp_path / "schedule.db")
    repository.import_schedule(ExcelScheduleParser().parse(schedule_workbook), sha256="original")
    workbook = load_workbook(schedule_workbook)
    for sheet in list(workbook):
        if sheet.title != "22 сентября 2026":
            workbook.remove(sheet)
    sheet = workbook.active
    sheet["A1"] = "ЗАНЯТИЙ НЕТ"
    for merged in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged))
    for row in sheet.iter_rows(min_row=106, max_row=109, min_col=3, max_col=18):
        for cell in row:
            cell.value = None
    path = tmp_path / "22.09.2026.xlsx"
    workbook.save(path)
    workbook.close()
    parsed = ExcelScheduleParser().parse(path)
    assert parsed.lessons == ()
    assert parsed.confirmed_empty_dates == (DAY,)
    repository.import_schedule(parsed, sha256="confirmed-empty")
    assert current(repository) == []


def test_empty_day_without_confirmation_keeps_previous_schedule(tmp_path: Path) -> None:
    repository = ScheduleRepository(tmp_path / "schedule.db")
    repository.import_schedule(schedule((lesson(),)), sha256="original")
    with pytest.raises(ImportQuarantinedError):
        repository.import_schedule(schedule(()), sha256="empty")
    assert current(repository) == [lesson()]


def test_mirrored_bell_table_is_not_a_missing_group(schedule_workbook: Path) -> None:
    workbook = load_workbook(schedule_workbook)
    sheet = workbook["22 сентября 2026"]
    sheet['S4'] = "Группа"
    sheet['T4'] = "День недели"
    sheet['S106'] = "1 пара"
    sheet['S107'] = "09.00-10.30"
    sheet['T106'] = "Вторник"
    workbook.save(schedule_workbook)
    workbook.close()
    assert ExcelScheduleParser().parse(schedule_workbook).validation_errors == ()


@pytest.mark.parametrize("blank_row", [6, 14])
def test_blank_pair_label_quarantines_small_day(tmp_path: Path, blank_row: int) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "22 сентября 2026"
    sheet['C4'] = "ИС-22"
    sheet['A6'] = "Вторник"
    for position, row in enumerate((6, 10, 14), 1):
        sheet.cell(row, 2, f"{position} пара")
        sheet.cell(row + 1, 2, "09.00-10.30")
        sheet.cell(row, 3, f"Предмет {position}")
        sheet.cell(row + 2, 3, "Иванов И.И.")
        sheet.cell(row + 1, 6, "44")
    path = tmp_path / "22.09.2026.xlsx"
    workbook.save(path)
    parser = ExcelScheduleParser()
    repository = ScheduleRepository(tmp_path / "schedule.db")
    original = parser.parse(path)
    assert len(original.lessons) == 3
    repository.import_schedule(original, sha256="original")
    sheet.cell(blank_row, 2).value = None
    workbook.save(path)
    workbook.close()
    partial = parser.parse(path)
    assert len(partial.lessons) == 2
    assert partial.validation_errors
    with pytest.raises(ImportQuarantinedError):
        repository.import_schedule(partial, sha256="blank-label")
    assert len(current(repository)) == 3


def test_footer_below_pairs_is_not_schedule_data(schedule_workbook: Path) -> None:
    workbook = load_workbook(schedule_workbook)
    sheet = workbook['22 сентября 2026']
    sheet['C111'] = 'Аудитории с 50 по 59 расположены в мастерских'
    sheet['A112'] = 'Заместитель директора по учебной работе'
    workbook.save(schedule_workbook)
    workbook.close()
    assert ExcelScheduleParser().parse(schedule_workbook).validation_errors == ()


class FakeBot:
    def __init__(self) -> None:
        self.sent: list[tuple[int, str]] = []

    async def send_message(self, chat_id: int, text: str, **kwargs) -> SimpleNamespace:
        self.sent.append((chat_id, text))
        return SimpleNamespace(message_id=len(self.sent))


@pytest.mark.asyncio
async def test_crash_after_import_resumes_outbox(tmp_path: Path) -> None:
    path = tmp_path / "schedule.db"
    repository = ScheduleRepository(path)
    repository.save_profile(123, role="student", target="ИС-22")
    repository.save_profile(456, role="teacher", target="Иванов И.И.")
    repository.save_chat_profile(-100, chat_type="supergroup", chat_title="ИС-22",
                                 target="ИС-22", configured_by=123)
    original = schedule((lesson(),))
    # A real abrupt exit leaves WAL and the committed outbox for the next process.
    code = """
import os, sys
from datetime import date
from schedule_bot.domain import Lesson, ParsedSchedule
from schedule_bot.repository import ScheduleRepository
day = date(2026, 9, 22)
lesson = Lesson(day, 'ИС-22', 1, '1 пара', 1, '09:00', '10:30',
                'Математика', 'Иванов И.И.', '44')
parsed = ParsedSchedule('22.09.2026.xlsx', day, day, (day,), (lesson,),
                        date_groups={day: ('ИС-22',)})
ScheduleRepository(sys.argv[1]).import_schedule(parsed, sha256='original')
os._exit(23)
"""
    environment = os.environ.copy()
    environment['PYTHONPATH'] = str(Path(__file__).resolve().parents[1] / 'src')
    crashed = subprocess.run(
        [sys.executable, '-B', '-c', code, str(path)], env=environment,
        capture_output=True, timeout=20,
    )
    assert crashed.returncode == 23, crashed.stderr.decode('utf-8', errors='replace')
    assert repository.notification_outbox_counts() == {"pending": 3}
    # No notifier call before the simulated process restart.
    restarted = ScheduleRepository(path)
    assert restarted.import_schedule(original, sha256="original").skipped
    bot = FakeBot()
    notifier = ScheduleNotifier(bot, restarted, (), broadcast_interval=0,
                                private_interval=0, group_interval=0)
    worker = asyncio.create_task(notifier.run_delivery_worker(idle_seconds=0.01))
    try:
        async def drained() -> None:
            while len(bot.sent) < 3:
                await asyncio.sleep(0.01)
        await asyncio.wait_for(drained(), timeout=2)
    finally:
        worker.cancel()
        with pytest.raises(asyncio.CancelledError):
            await worker
    assert {destination for destination, _ in bot.sent} == {123, 456, -100}
    assert restarted.notification_outbox_counts() == {"sent": 3}
    await notifier.notify_import(restarted.import_schedule(original, sha256="original"))
    assert len(bot.sent) == 3


def test_outbox_insert_failure_rolls_back_entire_import(tmp_path: Path) -> None:
    repository = ScheduleRepository(tmp_path / "schedule.db")
    repository.import_schedule(schedule((lesson(),)), sha256="original")
    repository.save_profile(123, role="student", target="ИС-22")
    with closing(sqlite3.connect(repository.database_path)) as connection, connection:
        connection.execute("CREATE TRIGGER fail_delivery BEFORE INSERT ON notification_outbox "
                           "BEGIN SELECT RAISE(ABORT, 'outbox failure'); END")
    with pytest.raises(sqlite3.IntegrityError, match="outbox failure"):
        repository.import_schedule(schedule((lesson(subject="Физика"),)), sha256="updated")
    assert current(repository)[0].subject == "Математика"
    assert repository.notification_outbox_counts() == {}
    with closing(sqlite3.connect(repository.database_path)) as connection, connection:
        assert connection.execute("SELECT COUNT(*) FROM imports").fetchone()[0] == 1
        connection.execute("DROP TRIGGER fail_delivery")
    assert not repository.import_schedule(schedule((lesson(subject="Физика"),)), sha256="updated").skipped
    assert repository.notification_outbox_counts() == {"pending": 1}


@pytest.mark.asyncio
async def test_queued_text_is_immutable_across_later_imports(tmp_path: Path) -> None:
    repository = ScheduleRepository(tmp_path / "schedule.db")
    repository.save_profile(123, role="student", target="ИС-22")
    repository.import_schedule(schedule((lesson(),)), sha256="original")
    repository.import_schedule(schedule((lesson(subject="Физика"),)), sha256="updated")
    bot = FakeBot()
    await ScheduleNotifier(bot, repository, (), broadcast_interval=0,
                           private_interval=0, group_interval=0).deliver_pending()
    assert len(bot.sent) == 2
    assert "Математика" in bot.sent[0][1] and "Физика" not in bot.sent[0][1]
    assert "Расписание изменилось" in bot.sent[1][1]


def test_subscriptions_are_captured_at_commit(tmp_path: Path) -> None:
    repository = ScheduleRepository(tmp_path / "schedule.db")
    repository.save_profile(1, role="student", target="ИС-22", notifications=False)
    repository.save_profile(2, role="student", target="ПД-12")
    repository.import_schedule(schedule((lesson(),)), sha256="original")
    repository.save_profile(3, role="student", target="ИС-22")
    assert repository.notification_outbox_counts() == {}


def test_concurrent_same_sha_creates_one_import_and_delivery(tmp_path: Path) -> None:
    repository = ScheduleRepository(tmp_path / "schedule.db")
    repository.save_profile(123, role="student", target="ИС-22")
    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: repository.import_schedule(schedule((lesson(),)), sha256="same"), range(2)))
    assert sorted(result.skipped for result in results) == [False, True]
    assert repository.notification_outbox_counts() == {"pending": 1}


def test_migrate_v4_preserves_data_outbox_and_backup(tmp_path: Path) -> None:
    path = tmp_path / "schedule.db"
    with closing(sqlite3.connect(path)) as connection, connection:
        for migration in (ScheduleRepository._migrate_to_v1, ScheduleRepository._migrate_to_v2,
                          ScheduleRepository._migrate_to_v3, ScheduleRepository._migrate_to_v4):
            migration(connection)
        connection.execute("INSERT INTO imports(id,source_filename,sha256,imported_at,start_date,end_date) "
                           "VALUES (1,'old.xlsx','old','2026-09-22','2026-09-22','2026-09-22')")
        connection.execute("INSERT INTO user_profiles VALUES (123,'student','ИС-22',0,'2026-09-22')")
        connection.execute("INSERT INTO notification_outbox(import_id,destination_id,destination_kind,chat_type,chunk_index,text,created_at) "
                           "VALUES (1,123,'user','private',0,'existing',?)",
                           (datetime.now(timezone.utc).isoformat(),))
        connection.execute("PRAGMA user_version = 4")
    repository = ScheduleRepository(path)
    backup = backup_before_migration(repository, tmp_path / "backups")
    assert backup is not None
    repository.initialize()
    assert repository.current_schema_version() == ScheduleRepository.SCHEMA_VERSION
    assert repository.get_profile(123)["notifications"] is False
    assert repository.pending_notifications()[0]["text"] == "existing"
    with closing(sqlite3.connect(backup)) as connection, connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 4
        assert connection.execute("SELECT text FROM notification_outbox").fetchone()[0] == "existing"


def test_cli_backs_up_before_schema_upgrade(tmp_path: Path) -> None:
    path = tmp_path / "schedule.db"
    with closing(sqlite3.connect(path)) as connection, connection:
        for migration in (
            ScheduleRepository._migrate_to_v1, ScheduleRepository._migrate_to_v2,
            ScheduleRepository._migrate_to_v3, ScheduleRepository._migrate_to_v4,
        ):
            migration(connection)
        connection.execute("PRAGMA user_version = 4")
    settings = SimpleNamespace(database_path=path, backups_path=tmp_path / "backups")
    repository = _prepare_repository(settings)
    assert repository.current_schema_version() == ScheduleRepository.SCHEMA_VERSION
    backups = list(settings.backups_path.glob("pre-migration-v4-to-v*-*.db"))
    assert len(backups) == 1
    with closing(sqlite3.connect(backups[0])) as connection, connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 4
    _prepare_repository(settings)
    assert len(list(settings.backups_path.glob("*.db"))) == 1
