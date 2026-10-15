from __future__ import annotations

from dataclasses import replace
from datetime import datetime, time, timedelta
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from schedule_bot.parser import ExcelScheduleParser, ScheduleParseError
from schedule_bot.parser.workbook import WorkbookLimits, WorkbookValues


def rewrite(path: Path, changes: dict[str, bytes]):
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts.update(changes)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def reference(path: Path):
    parser = ExcelScheduleParser()
    workbook = load_workbook(path, data_only=True)
    try:

        class Adapter:
            sheets = [(ws.title, ws.title) for ws in workbook.worksheets]

            def read_sheet(self, title, part):
                return workbook[part]

        from schedule_bot.parser.excel import parse_filename_date_range

        start, end = parse_filename_date_range(path.name)
        return parser._parse_workbook(Adapter(), path.name, start, end)
    finally:
        workbook.close()


@pytest.mark.parametrize("mac_epoch", [False, True])
def test_value_reader_preserves_dates_numbers_merges_and_warnings(
    schedule_workbook, mac_epoch
):
    wb = load_workbook(schedule_workbook)
    if mac_epoch:
        wb.epoch = CALENDAR_MAC_1904
    wb.create_sheet("Объявления")
    ws = wb["22 сентября 2026"]
    ws["F107"] = datetime(2026, 9, 22, 10, 30)
    ws["J107"] = time(11, 45)
    ws["N107"] = timedelta(days=2, hours=3)
    ws["R107"] = True
    wb.save(schedule_workbook)
    wb.close()
    assert ExcelScheduleParser().parse(schedule_workbook) == reference(
        schedule_workbook
    )


def test_hidden_merged_values_and_missing_dimensions_match_openpyxl(schedule_workbook):
    with ZipFile(schedule_workbook) as archive:
        xml = ET.fromstring(archive.read("xl/worksheets/sheet2.xml"))
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    dimension = xml.find(ns + "dimension")
    xml.remove(dimension)
    row = next(r for r in xml.iter(ns + "row") if r.get("r") == "106")
    hidden = ET.SubElement(row, ns + "c", {"r": "D106", "t": "inlineStr"})
    ET.SubElement(ET.SubElement(hidden, ns + "is"), ns + "t").text = "Hidden by merge"
    rewrite(schedule_workbook, {"xl/worksheets/sheet2.xml": ET.tostring(xml)})
    parsed = ExcelScheduleParser().parse(schedule_workbook)
    assert parsed == reference(schedule_workbook)
    assert all("Hidden by merge" not in item.subject for item in parsed.lessons)


def test_cached_formula_and_shared_rich_strings_match_openpyxl(schedule_workbook):
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(schedule_workbook) as archive:
        xml = ET.fromstring(archive.read("xl/worksheets/sheet2.xml"))
        types = ET.fromstring(archive.read("[Content_Types].xml"))
    cell = next(c for c in xml.iter(ns + "c") if c.get("r") == "C106")
    cell.clear()
    cell.attrib.update(r="C106", t="s")
    ET.SubElement(cell, ns + "v").text = "0"
    room = next(c for c in xml.iter(ns + "c") if c.get("r") == "F107")
    room.clear()
    room.set("r", "F107")
    ET.SubElement(room, ns + "f").text = "40+3"
    ET.SubElement(room, ns + "v").text = "43"
    ET.SubElement(
        types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
        {
            "PartName": "/xl/sharedStrings.xml",
            "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
        },
    )
    rewrite(
        schedule_workbook,
        {
            "xl/worksheets/sheet2.xml": ET.tostring(xml),
            "[Content_Types].xml": ET.tostring(types),
            "xl/sharedStrings.xml": (
                '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                "<si><r><t>Теория </t></r><r><t>права</t></r>"
                '<rPh sb="0" eb="1"><t>phonetic hint</t></rPh></si></sst>'
            ).encode(),
        },
    )
    assert ExcelScheduleParser().parse(schedule_workbook) == reference(
        schedule_workbook
    )


@pytest.mark.parametrize(
    "limit,value",
    [
        ("max_archive_bytes", 100),
        ("max_members", 2),
        ("max_uncompressed_bytes", 100),
        ("max_part_bytes", 100),
        ("max_metadata_bytes", 100),
        ("max_sheets", 1),
        ("max_rows", 20),
        ("max_columns", 5),
        ("max_cells", 1),
        ("max_value_cells", 1),
        ("max_merges", 1),
        ("max_text_chars", 1),
    ],
)
def test_resource_budgets_reject_whole_workbook_instead_of_truncating(
    schedule_workbook, limit, value
):
    limits = replace(WorkbookLimits(), **{limit: value})
    with pytest.raises(ScheduleParseError):
        ExcelScheduleParser(limits=limits).parse(schedule_workbook)


def test_unused_sheet_xml_is_not_loaded(schedule_workbook):
    rewrite(schedule_workbook, {"xl/worksheets/sheet5.xml": b"not XML"})
    assert ExcelScheduleParser().parse(schedule_workbook).lessons


def test_missing_cells_are_not_retained_and_large_formatting_is_rejected(
    schedule_workbook,
):
    with WorkbookValues(schedule_workbook, WorkbookLimits()) as reader:
        sheet = reader.read_sheet(*reader.sheets[0])
        count = len(sheet.values)
        for row in range(200, 1000):
            assert sheet.cell(row, 200).value is None
        assert len(sheet.values) == count
    wb = load_workbook(schedule_workbook)
    wb.active.cell(20, 2000).number_format = "0"
    wb.save(schedule_workbook)
    wb.close()
    with pytest.raises(ScheduleParseError, match="dimensions"):
        ExcelScheduleParser().parse(schedule_workbook)


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_xml_entities_are_rejected_before_expansion(schedule_workbook, encoding):
    xml = (
        '<?xml version="1.0" encoding="'
        + encoding
        + '"?><!DOCTYPE x [<!ENTITY a "expanded">]><x>&a;</x>'
    )
    rewrite(schedule_workbook, {"xl/workbook.xml": xml.encode(encoding)})
    with pytest.raises(ScheduleParseError, match="declarations"):
        ExcelScheduleParser().parse(schedule_workbook)


def test_duplicate_zip_parts_are_rejected(schedule_workbook):
    with (
        ZipFile(schedule_workbook, "a") as archive,
        pytest.warns(UserWarning, match="Duplicate"),
    ):
        archive.writestr("xl/workbook.xml", b"duplicate")
    with pytest.raises(ScheduleParseError, match="duplicate"):
        ExcelScheduleParser().parse(schedule_workbook)
