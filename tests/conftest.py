from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet


GROUPS = {
    3: "ПД-12",
    7: "ПД-22",
    11: "ИС-22",
    15: "Т -21",
}


def _prepare_sheet(worksheet: Worksheet) -> None:
    worksheet["A4"] = "День недели"
    worksheet["B4"] = "Группа"
    for column, group in GROUPS.items():
        worksheet.cell(4, column, group)


def _write_lesson(
    worksheet: Worksheet,
    *,
    day_row: int,
    weekday: str,
    time_range: str,
    entries: dict[int, tuple[str, str, str]],
) -> None:
    worksheet.cell(day_row, 1, weekday)
    worksheet.cell(day_row, 2, "1 пара")
    worksheet.cell(day_row + 1, 2, time_range)
    for column, (subject, teacher, room) in entries.items():
        worksheet.cell(day_row, column, subject)
        worksheet.merge_cells(
            start_row=day_row,
            start_column=column,
            end_row=day_row + 1,
            end_column=column + 2,
        )
        worksheet.cell(day_row + 1, column + 3, room)
        worksheet.cell(day_row + 2, column, teacher)
        worksheet.merge_cells(
            start_row=day_row + 2,
            start_column=column,
            end_row=day_row + 3,
            end_column=column + 2,
        )


def build_schedule_workbook(path: Path) -> Path:
    workbook = Workbook()
    monday = workbook.active
    monday.title = "21 сентября 2026"
    tuesday = workbook.create_sheet("22 сентября 2026")
    wednesday = workbook.create_sheet("23 сентября 2026")
    before_range = workbook.create_sheet("18 сентября 2026")
    after_range = workbook.create_sheet("24 сентября 2026")

    for worksheet in workbook.worksheets:
        _prepare_sheet(worksheet)

    _write_lesson(
        monday,
        day_row=6,
        weekday="Понедельник",
        time_range="09.00-10.10",
        entries={3: ("МДК.01.03.", "Яцишина Н.В.", "44")},
    )
    _write_lesson(
        monday,
        day_row=30,
        weekday="Пятница",
        time_range="09.00-10.30",
        entries={3: ("Экономика организ.", "Другой преподаватель", "31")},
    )
    _write_lesson(
        tuesday,
        day_row=106,
        weekday="Вторник",
        time_range="09.00-10.30",
        entries={
            3: ("Теория гос. и права", "Булаева В.В.", "43"),
            7: ("Криминалистика", "Яцишина Н.В.", "28"),
            11: ("МДК.01.04.", "Паталах Ю.А.", "45"),
            15: ("Физкультура", "Крайнов М.В.", "спортзал"),
        },
    )
    _write_lesson(
        wednesday,
        day_row=106,
        weekday="Среда",
        time_range="09.00-10.30",
        entries={3: ("Психология", "Коваленко М.В.", "23")},
    )
    _write_lesson(
        before_range,
        day_row=106,
        weekday="Пятница",
        time_range="09.00-10.30",
        entries={3: ("Старое расписание", "Преподаватель", "1")},
    )
    _write_lesson(
        after_range,
        day_row=106,
        weekday="Четверг",
        time_range="09.00-10.30",
        entries={3: ("Будущее расписание", "Преподаватель", "2")},
    )
    before_range.sheet_state = "hidden"
    after_range.sheet_state = "hidden"
    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture
def schedule_workbook(tmp_path: Path) -> Path:
    return build_schedule_workbook(tmp_path / "21.-23.09.2026.xlsx")

